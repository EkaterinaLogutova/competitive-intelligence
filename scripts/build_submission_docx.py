from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "submission"
OUTPUT_DIR = SOURCE_DIR / "final"
GITHUB_URL = "https://github.com/EkaterinaLogutova/competitive-intelligence"
AUTHOR = "Екатерина Логутова"
TEAM = (
    "Команда исследований технологий управления интеллектуальной собственности "
    "на основе ИИ и аналитики"
)

OUTPUT_NAMES = {
    "task-1.md": (
        "Логутова Е. Команда исследований технологий управления интеллектуальной "
        "собственностью на основе ИИ и аналитики. Задание 1.docx"
    ),
    "task-2.md": (
        "Логутова Е. Команда исследований технологий управления интеллектуальной "
        "собственностью на основе ИИ и аналитики. Задание 2.docx"
    ),
    "task-3.md": (
        "Логутова Е. Команда исследований технологий управления интеллектуальной "
        "собственностью на основе ИИ и аналитики. Задание 3.docx"
    ),
}


def set_run_font(run, name: str, size: float, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def set_cell_margins(cell, top: int = 60, start: int = 75, bottom: int = 60, end: int = 75) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_borders(table) -> None:
    table_pr = table._tbl.tblPr
    borders = table_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)


def remove_paragraph_borders(paragraph_or_style) -> None:
    paragraph_properties = paragraph_or_style._element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color, underline))
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("`", "").replace("**", "")


def configure_document(document: Document, title: str, compact: bool = False) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.6 if compact else 1.8)
    section.bottom_margin = Cm(1.6 if compact else 1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.7 if compact else 10)
    normal.paragraph_format.space_after = Pt(4 if compact else 5)
    normal.paragraph_format.line_spacing = 1.05 if compact else 1.1

    for style_name, size, space_before, space_after in (
        ("Title", 17, 0, 12),
        ("Heading 1", 13, 14, 6),
        ("Heading 2", 11.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(
            max(0, space_before - 2) if compact else space_before
        )
        style.paragraph_format.space_after = Pt(
            max(0, space_after - 1) if compact else space_after
        )
        style.paragraph_format.keep_with_next = True

    remove_paragraph_borders(styles["Title"])

    document.core_properties.author = AUTHOR
    document.core_properties.title = title
    document.core_properties.subject = TEAM


def add_identity_block(document: Document, title: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    remove_paragraph_borders(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(title)
    set_run_font(run, "Arial", 17, bold=True)

    author_paragraph = document.add_paragraph()
    author_run = author_paragraph.add_run(AUTHOR)
    set_run_font(author_run, "Arial", 11, bold=True)

    team_paragraph = document.add_paragraph()
    team_run = team_paragraph.add_run(TEAM)
    set_run_font(team_run, "Arial", 10.5)

    repository_paragraph = document.add_paragraph()
    label_run = repository_paragraph.add_run("Репозиторий проекта: ")
    set_run_font(label_run, "Arial", 10.5)
    add_hyperlink(repository_paragraph, GITHUB_URL, GITHUB_URL)
    repository_paragraph.paragraph_format.space_after = Pt(12)


def add_table(document: Document, rows: list[list[str]]) -> None:
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)

    for row_index, values in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            value = clean_inline(values[column_index]) if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            set_run_font(run, "Arial", 8, bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, "E7E6E6")

    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        values = [value.strip() for value in line.strip().strip("|").split("|")]
        if all(value and set(value) <= {"-", ":"} for value in values):
            continue
        rows.append(values)
    return rows


def add_body(document: Document, markdown: str, compact: bool = False) -> None:
    lines = markdown.splitlines()
    index = 0
    paragraph_buffer: list[str] = []
    in_code = False
    code_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        text = clean_inline(" ".join(part.strip() for part in paragraph_buffer))
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.keep_together = False
        paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.space_after = Pt(8)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, "Consolas", 9)
                code_lines.clear()
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            document.add_paragraph(clean_inline(stripped[3:]), style="Heading 1")
            index += 1
            continue

        if stripped.startswith("### "):
            flush_paragraph()
            document.add_paragraph(clean_inline(stripped[4:]), style="Heading 2")
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, parse_table(table_lines))
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(clean_inline(stripped[2:]))
            set_run_font(run, "Arial", 9.7 if compact else 10.5)
            index += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.65)
            paragraph.paragraph_format.first_line_indent = Cm(-0.45)
            run = paragraph.add_run(clean_inline(stripped))
            set_run_font(run, "Arial", 9.7 if compact else 10)
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()


def build_document(source_path: Path, output_path: Path) -> None:
    markdown = source_path.read_text(encoding="utf-8")
    first_line = next(line for line in markdown.splitlines() if line.startswith("# "))
    title = first_line.removeprefix("# ").strip()

    document = Document()
    compact = source_path.name == "task-3.md"
    configure_document(document, title, compact=compact)
    add_identity_block(document, title)
    add_body(document, markdown, compact=compact)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for source_name, output_name in OUTPUT_NAMES.items():
        build_document(SOURCE_DIR / source_name, OUTPUT_DIR / output_name)
        print(OUTPUT_DIR / output_name)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
